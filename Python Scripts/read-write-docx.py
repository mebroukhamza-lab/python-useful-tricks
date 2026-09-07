# pip install python-docx
from docx import Document
doc = Document()
doc.add_heading("My Report", level=1)
doc.add_paragraph("This is an automatically generated paragraph.")
doc.save("report.docx")
print("Word document created successfully")
